import csv
import io
from datetime import datetime, date, timedelta
from typing import List, Optional, Union

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import load_workbook, Workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError

from database import engine, get_db, Base
from models import VideoMaterial, AdData, ROISnapshot
from schemas import (
    VideoMaterialCreate, VideoMaterialUpdate, VideoMaterialOut,
    AdDataCreate, AdDataUpdate, AdDataOut,
    AdDataAggregatedOut, AdDataImportResult, ImportFailureDetail,
    ROISnapshotCreate, ROISnapshotUpdate, ROISnapshotOut,
    ROIDataOut, ROIBatchItemOut,
    MaterialROIItemOut, MaterialROIFilterResult,
    ROIThresholdsOut, ROIThresholdsUpdate,
    ROITrendDataPoint, ROITrendResponse, ROICompareVideoItem, ROICompareResponse,
    HealthResponse,
)
import crud
from roi_engine import trigger_roi_recalculation, calculate_video_roi, calculate_roi, get_roi_label

Base.metadata.create_all(bind=engine)

app = FastAPI(title="视频号微信豆投放素材筛选系统", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/api/health", response_model=HealthResponse)
def health_check():
    return HealthResponse()


@app.post("/api/materials", response_model=VideoMaterialOut)
def create_material(data: VideoMaterialCreate, db: Session = Depends(get_db)):
    return crud.create_video_material(db, data)


@app.get("/api/materials", response_model=List[VideoMaterialOut])
def list_materials(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return crud.get_list(db, VideoMaterial, skip, limit)


@app.get("/api/materials/roi-batch", response_model=List[ROIBatchItemOut])
def batch_get_roi(
    video_ids: str = Query(..., description="逗号分隔的视频ID"),
    date_range_type: str = Query('all'),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    ids = [int(x.strip()) for x in video_ids.split(",") if x.strip()]
    results = []
    for vid in ids:
        video = crud.get_by_id(db, VideoMaterial, vid)
        if not video:
            continue
        roi_data = calculate_video_roi(db, vid, date_range_type=date_range_type, start_date=start_date, end_date=end_date)
        results.append(ROIBatchItemOut(video_id=vid, **roi_data))
    return results


_roi_thresholds = {"high_threshold": 100.0, "low_threshold": 0.0}


@app.get("/api/materials/roi-filter", response_model=MaterialROIFilterResult)
def roi_filter_materials(
    roi_min: Optional[float] = Query(None),
    roi_max: Optional[float] = Query(None),
    sort_by: str = Query("roi"),
    sort_order: str = Query("desc"),
    material_type: Optional[str] = Query(None),
    duration_min: Optional[float] = Query(None),
    duration_max: Optional[float] = Query(None),
    date_range_type: str = Query("all"),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    status: Optional[str] = Query(None),
    skip: int = Query(0),
    limit: int = Query(20),
    db: Session = Depends(get_db),
):
    q = db.query(VideoMaterial)
    if material_type:
        q = q.filter(VideoMaterial.material_type == material_type)
    if duration_min is not None:
        q = q.filter(VideoMaterial.duration >= duration_min)
    if duration_max is not None:
        q = q.filter(VideoMaterial.duration <= duration_max)
    if status:
        q = q.filter(VideoMaterial.status == status)

    materials = q.all()

    items = []
    for m in materials:
        roi_data = calculate_video_roi(db, m.id, date_range_type=date_range_type, start_date=start_date, end_date=end_date)
        roi_val = roi_data.get("roi")

        if roi_min is not None and (roi_val is None or roi_val < roi_min):
            continue
        if roi_max is not None and (roi_val is None or roi_val > roi_max):
            continue

        label = get_roi_label(roi_val, _roi_thresholds["high_threshold"], _roi_thresholds["low_threshold"])

        items.append(MaterialROIItemOut(
            id=m.id,
            title=m.title,
            video_url=m.video_url,
            cover_url=m.cover_url,
            duration=m.duration,
            aspect_ratio=m.aspect_ratio,
            material_type=m.material_type,
            status=m.status,
            created_at=m.created_at,
            updated_at=m.updated_at,
            roi=roi_val,
            bean_output=roi_data.get("bean_output"),
            ctr=roi_data.get("ctr"),
            interaction_rate=roi_data.get("interaction_rate"),
            conversion_rate=roi_data.get("conversion_rate"),
            total_bean_cost=roi_data.get("total_bean_cost", 0.0),
            total_gmv=roi_data.get("total_gmv", 0.0),
            roi_label=label,
        ))

    sort_field_map = {
        "roi": "roi",
        "bean_output": "bean_output",
        "ctr": "ctr",
        "interaction_rate": "interaction_rate",
        "conversion_rate": "conversion_rate",
        "created_at": "created_at",
    }
    field = sort_field_map.get(sort_by, "roi")
    reverse = sort_order == "desc"

    def sort_key(item):
        val = getattr(item, field, None)
        if val is None:
            return (1, 0)
        return (0, val)

    items.sort(key=sort_key, reverse=reverse)

    total = len(items)
    paged = items[skip: skip + limit]

    return MaterialROIFilterResult(total=total, items=paged)


@app.get("/api/materials/export")
def export_materials(
    roi_min: Optional[float] = Query(None),
    roi_max: Optional[float] = Query(None),
    sort_by: str = Query("roi"),
    sort_order: str = Query("desc"),
    material_type: Optional[str] = Query(None),
    duration_min: Optional[float] = Query(None),
    duration_max: Optional[float] = Query(None),
    date_range_type: str = Query("all"),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    q = db.query(VideoMaterial)
    if material_type:
        q = q.filter(VideoMaterial.material_type == material_type)
    if duration_min is not None:
        q = q.filter(VideoMaterial.duration >= duration_min)
    if duration_max is not None:
        q = q.filter(VideoMaterial.duration <= duration_max)
    if status:
        q = q.filter(VideoMaterial.status == status)

    materials = q.all()

    items = []
    for m in materials:
        roi_data = calculate_video_roi(db, m.id, date_range_type=date_range_type, start_date=start_date, end_date=end_date)
        roi_val = roi_data.get("roi")

        if roi_min is not None and (roi_val is None or roi_val < roi_min):
            continue
        if roi_max is not None and (roi_val is None or roi_val > roi_max):
            continue

        label = get_roi_label(roi_val, _roi_thresholds["high_threshold"], _roi_thresholds["low_threshold"])

        items.append({
            "title": m.title,
            "material_type": m.material_type,
            "duration": m.duration,
            "roi": roi_val,
            "bean_output": roi_data.get("bean_output"),
            "ctr": roi_data.get("ctr"),
            "interaction_rate": roi_data.get("interaction_rate"),
            "conversion_rate": roi_data.get("conversion_rate"),
            "total_bean_cost": roi_data.get("total_bean_cost", 0.0),
            "total_gmv": roi_data.get("total_gmv", 0.0),
            "roi_label": label,
            "status": m.status,
            "created_at": m.created_at,
        })

    sort_field_map = {
        "roi": "roi",
        "bean_output": "bean_output",
        "ctr": "ctr",
        "interaction_rate": "interaction_rate",
        "conversion_rate": "conversion_rate",
        "created_at": "created_at",
    }
    field = sort_field_map.get(sort_by, "roi")
    reverse = sort_order == "desc"

    def sort_key(item):
        val = item.get(field)
        if val is None:
            return (1, 0)
        return (0, val)

    items.sort(key=sort_key, reverse=reverse)

    wb = Workbook()
    ws = wb.active
    ws.title = "素材ROI数据"

    headers = ["序号", "标题", "类型", "时长(秒)", "ROI(%)", "单豆产出", "点击率(%)", "互动率(%)", "转化率(%)", "微信豆消耗", "GMV", "ROI标签", "状态"]
    ws.append(headers)

    status_map = {"active": "正常", "pending_elimination": "待淘汰"}
    label_map = {"high": "高ROI", "low": "低ROI"}
    type_map = {"video": "视频", "image": "图片"}

    for idx, item in enumerate(items, start=1):
        ws.append([
            idx,
            item["title"],
            type_map.get(item["material_type"], item["material_type"] or ""),
            item["duration"],
            item["roi"],
            item["bean_output"],
            item["ctr"],
            item["interaction_rate"],
            item["conversion_rate"],
            item["total_bean_cost"],
            item["total_gmv"],
            label_map.get(item["roi_label"], "") if item["roi_label"] else "",
            status_map.get(item["status"], item["status"] or ""),
        ])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    today_str = date.today().strftime("%Y%m%d")
    filename = f"materials_roi_export_{today_str}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.get("/api/settings/roi-thresholds", response_model=ROIThresholdsOut)
def get_roi_thresholds():
    return ROIThresholdsOut(**_roi_thresholds)


@app.put("/api/settings/roi-thresholds", response_model=ROIThresholdsOut)
def update_roi_thresholds(data: ROIThresholdsUpdate):
    if data.high_threshold is not None:
        _roi_thresholds["high_threshold"] = data.high_threshold
    if data.low_threshold is not None:
        _roi_thresholds["low_threshold"] = data.low_threshold
    return ROIThresholdsOut(**_roi_thresholds)


@app.get("/api/materials/roi-compare", response_model=ROICompareResponse)
def roi_compare(
    video_ids: str = Query(..., description="逗号分隔的视频ID，最多5个"),
    group_by: str = Query("day", description="分组方式: day/week"),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    ids = [int(x.strip()) for x in video_ids.split(",") if x.strip()]
    if len(ids) > 5:
        raise HTTPException(status_code=400, detail="最多支持5个视频对比")
    if not ids:
        raise HTTPException(status_code=400, detail="请至少选择1个视频")

    today = date.today()
    if start_date is None:
        start_date = today - timedelta(days=30)
    if end_date is None:
        end_date = today

    start_datetime = datetime.combine(start_date, datetime.min.time())
    end_datetime = datetime.combine(end_date, datetime.max.time())

    videos_data = []
    for vid in ids:
        video = crud.get_by_id(db, VideoMaterial, vid)
        if not video:
            continue
        aggregated = crud.aggregate_ad_data(
            db, video_id=vid, start_date=start_datetime, end_date=end_datetime, group_by=group_by
        )
        data_points = []
        for row in aggregated:
            metrics = calculate_roi(
                bean_cost=row["bean_cost"],
                gmv=row["gmv"],
                impressions=row["impressions"],
                clicks=row["clicks"],
                interactions=row["interactions"],
                conversions=row["conversions"],
            )
            data_points.append(ROITrendDataPoint(
                period=row["period"],
                beanCost=row["bean_cost"],
                gmv=row["gmv"],
                roi=metrics["roi"],
                ctr=metrics["ctr"],
                interactionRate=metrics["interaction_rate"],
                conversionRate=metrics["conversion_rate"],
            ))
        videos_data.append(ROICompareVideoItem(
            videoId=vid,
            videoTitle=video.title,
            data=data_points,
        ))

    return ROICompareResponse(groupBy=group_by, videos=videos_data)


@app.get("/api/materials/{material_id}", response_model=VideoMaterialOut)
def get_material(material_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, VideoMaterial, material_id)
    if not obj:
        raise HTTPException(status_code=404, detail="素材不存在")
    return obj


@app.put("/api/materials/{material_id}", response_model=VideoMaterialOut)
def update_material(material_id: int, data: VideoMaterialUpdate, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, VideoMaterial, material_id)
    if not obj:
        raise HTTPException(status_code=404, detail="素材不存在")
    return crud.update_video_material(db, obj, data)


@app.delete("/api/materials/{material_id}")
def delete_material(material_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, VideoMaterial, material_id)
    if not obj:
        raise HTTPException(status_code=404, detail="素材不存在")
    crud.delete_video_material(db, obj)
    return {"ok": True}


@app.put("/api/materials/{material_id}/eliminate", response_model=VideoMaterialOut)
def eliminate_material(material_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, VideoMaterial, material_id)
    if not obj:
        raise HTTPException(status_code=404, detail="素材不存在")
    obj.status = "pending_elimination"
    db.commit()
    db.refresh(obj)
    return obj


@app.put("/api/materials/{material_id}/restore", response_model=VideoMaterialOut)
def restore_material(material_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, VideoMaterial, material_id)
    if not obj:
        raise HTTPException(status_code=404, detail="素材不存在")
    obj.status = "active"
    db.commit()
    db.refresh(obj)
    return obj


@app.post("/api/ad-data", response_model=AdDataOut)
def create_ad_data(data: AdDataCreate, db: Session = Depends(get_db)):
    video = crud.get_by_id(db, VideoMaterial, data.video_id)
    if not video:
        raise HTTPException(status_code=400, detail="视频ID不存在")
    obj = crud.create_ad_data(db, data)
    trigger_roi_recalculation(db, obj.video_id)
    return obj


@app.get("/api/ad-data", response_model=Union[List[AdDataOut], List[AdDataAggregatedOut]])
def list_ad_data(
    video_id: Optional[int] = Query(None),
    start_date: Optional[datetime] = Query(None),
    end_date: Optional[datetime] = Query(None),
    group_by: Optional[str] = Query(None),
    skip: int = Query(0),
    limit: int = Query(100),
    db: Session = Depends(get_db),
):
    if group_by:
        return crud.aggregate_ad_data(db, video_id=video_id, start_date=start_date, end_date=end_date, group_by=group_by)
    return crud.query_ad_data(db, video_id=video_id, start_date=start_date, end_date=end_date, skip=skip, limit=limit)


@app.get("/api/ad-data/{ad_data_id}", response_model=AdDataOut)
def get_ad_data(ad_data_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, AdData, ad_data_id)
    if not obj:
        raise HTTPException(status_code=404, detail="投放数据不存在")
    return obj


@app.put("/api/ad-data/{ad_data_id}", response_model=AdDataOut)
def update_ad_data(ad_data_id: int, data: AdDataUpdate, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, AdData, ad_data_id)
    if not obj:
        raise HTTPException(status_code=404, detail="投放数据不存在")
    updated = crud.update_ad_data(db, obj, data)
    trigger_roi_recalculation(db, updated.video_id)
    return updated


@app.delete("/api/ad-data/{ad_data_id}")
def delete_ad_data(ad_data_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, AdData, ad_data_id)
    if not obj:
        raise HTTPException(status_code=404, detail="投放数据不存在")
    crud.delete_ad_data(db, obj)
    return {"ok": True}


@app.post("/api/ad-data/import", response_model=AdDataImportResult)
def import_ad_data(file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = file.filename or ""
    content = file.file.read()

    if filename.endswith(".xlsx"):
        rows = _parse_xlsx(content)
    elif filename.endswith(".csv"):
        rows = _parse_csv(content)
    else:
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 和 .csv 格式")

    success_count = 0
    failures: List[ImportFailureDetail] = []
    create_list: List[AdDataCreate] = []
    affected_video_ids = set()

    for idx, row in enumerate(rows, start=2):
        try:
            title = str(row.get("视频标题", "")).strip()
            date_str = str(row.get("日期", "")).strip()
            bean_cost = float(row.get("微信豆消耗", 0) or 0)
            impressions = int(float(row.get("曝光量", 0) or 0))
            clicks = int(float(row.get("点击量", 0) or 0))
            interactions = int(float(row.get("互动量", 0) or 0))
            conversions = int(float(row.get("转化数", 0) or 0))
            gmv = float(row.get("GMV", 0) or 0)

            if not title:
                failures.append(ImportFailureDetail(row=idx, reason="视频标题为空"))
                continue

            video = crud.get_video_by_title(db, title)
            if not video:
                failures.append(ImportFailureDetail(row=idx, reason=f"视频标题「{title}」未找到匹配"))
                continue

            date_val = datetime.strptime(date_str, "%Y-%m-%d")

            create_list.append(AdDataCreate(
                video_id=video.id,
                date=date_val,
                bean_cost=bean_cost,
                impressions=impressions,
                clicks=clicks,
                interactions=interactions,
                conversions=conversions,
                gmv=gmv,
            ))
            affected_video_ids.add(video.id)
        except (ValueError, KeyError) as e:
            failures.append(ImportFailureDetail(row=idx, reason=str(e)))

    if create_list:
        try:
            crud.batch_create_ad_data(db, create_list)
            success_count = len(create_list)
        except IntegrityError:
            db.rollback()
            raise HTTPException(status_code=500, detail="批量导入数据库错误")

    for vid in affected_video_ids:
        trigger_roi_recalculation(db, vid)

    return AdDataImportResult(
        success_count=success_count,
        failure_count=len(failures),
        failures=failures,
    )


@app.get("/api/materials/{video_id}/ad-data", response_model=Union[List[AdDataOut], List[AdDataAggregatedOut]])
def list_ad_data_by_video(
    video_id: int,
    start_date: Optional[datetime] = Query(None),
    end_date: Optional[datetime] = Query(None),
    group_by: Optional[str] = Query(None),
    skip: int = Query(0),
    limit: int = Query(100),
    db: Session = Depends(get_db),
):
    video = crud.get_by_id(db, VideoMaterial, video_id)
    if not video:
        raise HTTPException(status_code=404, detail="视频不存在")
    if group_by:
        return crud.aggregate_ad_data(db, video_id=video_id, start_date=start_date, end_date=end_date, group_by=group_by)
    return crud.query_ad_data(db, video_id=video_id, start_date=start_date, end_date=end_date, skip=skip, limit=limit)


@app.post("/api/roi-snapshots", response_model=ROISnapshotOut)
def create_roi_snapshot(data: ROISnapshotCreate, db: Session = Depends(get_db)):
    return crud.create_roi_snapshot(db, data)


@app.get("/api/roi-snapshots", response_model=List[ROISnapshotOut])
def list_roi_snapshots(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return crud.get_list(db, ROISnapshot, skip, limit)


@app.get("/api/roi-snapshots/{snapshot_id}", response_model=ROISnapshotOut)
def get_roi_snapshot(snapshot_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, ROISnapshot, snapshot_id)
    if not obj:
        raise HTTPException(status_code=404, detail="ROI快照不存在")
    return obj


@app.put("/api/roi-snapshots/{snapshot_id}", response_model=ROISnapshotOut)
def update_roi_snapshot(snapshot_id: int, data: ROISnapshotUpdate, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, ROISnapshot, snapshot_id)
    if not obj:
        raise HTTPException(status_code=404, detail="ROI快照不存在")
    return crud.update_roi_snapshot(db, obj, data)


@app.delete("/api/roi-snapshots/{snapshot_id}")
def delete_roi_snapshot(snapshot_id: int, db: Session = Depends(get_db)):
    obj = crud.get_by_id(db, ROISnapshot, snapshot_id)
    if not obj:
        raise HTTPException(status_code=404, detail="ROI快照不存在")
    crud.delete_roi_snapshot(db, obj)
    return {"ok": True}


@app.get("/api/materials/{video_id}/roi-snapshots", response_model=List[ROISnapshotOut])
def list_roi_snapshots_by_video(video_id: int, skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    return crud.get_roi_snapshots_by_video(db, video_id, skip, limit)


@app.get("/api/materials/{video_id}/roi", response_model=ROIDataOut)
def get_video_roi(
    video_id: int,
    date_range_type: str = Query('all'),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    video = crud.get_by_id(db, VideoMaterial, video_id)
    if not video:
        raise HTTPException(status_code=404, detail="视频不存在")
    return calculate_video_roi(db, video_id, date_range_type=date_range_type, start_date=start_date, end_date=end_date)


@app.get("/api/materials/{video_id}/roi-trend", response_model=ROITrendResponse)
def get_video_roi_trend(
    video_id: int,
    group_by: str = Query("day", description="分组方式: day/week"),
    start_date: Optional[date] = Query(None),
    end_date: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    video = crud.get_by_id(db, VideoMaterial, video_id)
    if not video:
        raise HTTPException(status_code=404, detail="视频不存在")

    today = date.today()
    if start_date is None:
        start_date = today - timedelta(days=30)
    if end_date is None:
        end_date = today

    start_datetime = datetime.combine(start_date, datetime.min.time())
    end_datetime = datetime.combine(end_date, datetime.max.time())

    aggregated = crud.aggregate_ad_data(
        db, video_id=video_id, start_date=start_datetime, end_date=end_datetime, group_by=group_by
    )

    data_points = []
    for row in aggregated:
        metrics = calculate_roi(
            bean_cost=row["bean_cost"],
            gmv=row["gmv"],
            impressions=row["impressions"],
            clicks=row["clicks"],
            interactions=row["interactions"],
            conversions=row["conversions"],
        )
        data_points.append(ROITrendDataPoint(
            period=row["period"],
            beanCost=row["bean_cost"],
            gmv=row["gmv"],
            roi=metrics["roi"],
            ctr=metrics["ctr"],
            interactionRate=metrics["interaction_rate"],
            conversionRate=metrics["conversion_rate"],
        ))

    return ROITrendResponse(
        videoId=video_id,
        videoTitle=video.title,
        groupBy=group_by,
        data=data_points,
    )


@app.post("/api/materials/{video_id}/recalculate-roi")
def recalculate_roi(video_id: int, db: Session = Depends(get_db)):
    video = crud.get_by_id(db, VideoMaterial, video_id)
    if not video:
        raise HTTPException(status_code=404, detail="视频不存在")
    trigger_roi_recalculation(db, video_id)
    return {"ok": True, "message": "ROI重算完成"}


def _parse_xlsx(content: bytes) -> List[dict]:
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter, None)
    if not headers:
        wb.close()
        return []
    header_map = [str(h).strip() if h else "" for h in headers]
    result = []
    for row in rows_iter:
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(header_map):
                row_dict[header_map[i]] = val
        result.append(row_dict)
    wb.close()
    return result


def _parse_csv(content: bytes) -> List[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [row for row in reader]
